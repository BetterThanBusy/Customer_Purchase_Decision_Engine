"""
src/workbook.py

Builds outputs/video_tables.xlsx — the tables meant to be shown on screen
rather than read in a terminal.

The break-even sheet carries LIVE Excel formulas. Change the margin or the
discount rate on camera and the whole grid recalculates, which is a far
better demonstration than a static screenshot.

Colour convention follows standard financial-model practice:
  blue text   = an input you are meant to change
  black text  = a formula
  yellow fill = key assumption
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.engine import (
    ACTION_CATALOGUE,
    STATE_PLAYBOOK,
    GROSS_MARGIN_RATE,
    CONTROL_HOLDOUT_SHARE,
)


# ==========================================================================
# STYLE
# ==========================================================================

FONT = "Arial"

INK = "000000"
BURGUNDY = "51051E"
BEIGE = "F7F1EA"
TAN = "E3DACF"
TEAL = "043C4D"
ORANGE = "FF4500"
BLUE = "0000FF"
GREEN = "043C4D"  # teal stands in for the cross-sheet link colour

HEADER_FILL = PatternFill("solid", fgColor=BURGUNDY)
BAND_FILL = PatternFill("solid", fgColor=BEIGE)
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFF00")
SECTION_FILL = PatternFill("solid", fgColor=TAN)

HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT, bold=True, color=BURGUNDY, size=14)
BODY_FONT = Font(name=FONT, size=10, color=INK)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="5A5A5A")
INPUT_FONT = Font(name=FONT, size=10, bold=True, color=BLUE)
FORMULA_FONT = Font(name=FONT, size=10, color=INK)
LINK_FONT = Font(name=FONT, size=10, color=GREEN)

THIN = Side(style="thin", color="C9BCAA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _title(sheet, text, subtitle=None):
    sheet["A1"] = text
    sheet["A1"].font = TITLE_FONT
    if subtitle:
        sheet["A2"] = subtitle
        sheet["A2"].font = NOTE_FONT
    sheet.sheet_view.showGridLines = False


def _write_table(sheet, frame, start_row, start_col=1, number_formats=None,
                 band=True, width_cap=42):
    """Write a dataframe as a formatted table. Returns the last row used."""
    number_formats = number_formats or {}

    for j, column in enumerate(frame.columns):
        cell = sheet.cell(row=start_row, column=start_col + j, value=str(column))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX

    for i, (_, row) in enumerate(frame.iterrows()):
        for j, column in enumerate(frame.columns):
            value = row[column]
            if isinstance(value, (np.integer,)):
                value = int(value)
            elif isinstance(value, (np.floating,)):
                value = None if pd.isna(value) else float(value)
            elif isinstance(value, pd.Timestamp):
                value = value.strftime("%Y-%m-%d")
            elif pd.isna(value) if np.isscalar(value) else False:
                value = None

            cell = sheet.cell(row=start_row + 1 + i, column=start_col + j, value=value)
            cell.font = BODY_FONT
            cell.border = BOX
            if band and i % 2 == 1:
                cell.fill = BAND_FILL
            if column in number_formats:
                cell.number_format = number_formats[column]

    for j, column in enumerate(frame.columns):
        letter = get_column_letter(start_col + j)
        longest = max(
            [len(str(column))] + [len(str(v)) for v in frame[column].head(60)]
        )
        sheet.column_dimensions[letter].width = min(max(11, longest + 3), width_cap)

    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)
    return start_row + len(frame)


def _note(sheet, row, text):
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = NOTE_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    return row + 1


# ==========================================================================
# SHEETS
# ==========================================================================

def _sheet_legend(workbook, best_name, test_metrics, observed_rate,
                  mean_balanced, mean_unweighted, median_aov):
    sheet = workbook.create_sheet("00 Read me")
    _title(
        sheet,
        "Customer Purchase Decision Engine — presentation tables",
        "Companion to outputs/charts/. Each sheet maps to a beat in the video.",
    )

    rows = [
        ("", ""),
        ("COLOUR CONVENTION", ""),
        ("Blue bold text", "an input you can change — the sheet recalculates"),
        ("Black text", "a formula or a computed result"),
        ("Yellow fill", "a key assumption, stated rather than buried"),
        ("", ""),
        ("SHEET MAP", ""),
        ("01 Customer snapshot", "what one row of the modelling table looks like"),
        ("02 Model scorecard", "baselines, model, calibration — the evidence it works"),
        ("03 Customer states", "the segmentation, with size, likelihood and value"),
        ("04 Action catalogue", "every action priced by what it costs, not what it is called"),
        ("05 Break-even calculator", "LIVE. Change margin or discount and watch the grid move"),
        ("06 State x action", "which state receives which action"),
        ("07 Priority customers", "the list a marketer would actually work"),
        ("08 Budget", "cost of this cycle and the uplift it must deliver"),
        ("09 Experiment design", "control cells and what they can detect"),
        ("10 Playbook", "what each segment actually gets — do, do not, and how it is judged"),
        ("", ""),
        ("HEADLINE NUMBERS", ""),
        ("Model selected", best_name),
        ("Test ROC-AUC", round(test_metrics["roc_auc"], 4)),
        ("Test PR-AUC", round(test_metrics["pr_auc"], 4)),
        ("Test base rate", round(test_metrics["base_rate"], 4)),
        ("Median AOV used in economics", round(median_aov, 2)),
        ("Gross margin assumption", GROSS_MARGIN_RATE),
        ("Control holdout share", CONTROL_HOLDOUT_SHARE),
        ("", ""),
        ("THE CALIBRATION TRAP", ""),
        ("Observed repeat rate (validation)", round(observed_rate, 4)),
        ("Mean predicted, class_weight='balanced'", round(mean_balanced, 4)),
        ("Mean predicted, unweighted", round(mean_unweighted, 4)),
    ]

    for i, (label, value) in enumerate(rows, start=4):
        left = sheet.cell(row=i, column=1, value=label)
        left.font = Font(name=FONT, size=10, bold=label.isupper() and label != "",
                         color=BURGUNDY if label.isupper() and label else INK)
        if label.isupper() and label:
            left.fill = SECTION_FILL
        right = sheet.cell(row=i, column=2, value=value)
        right.font = BODY_FONT
        if isinstance(value, float) and value <= 1:
            right.number_format = "0.0%" if label.endswith(("share", "assumption", "rate")) else "0.0000"

    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 48

    _note(
        sheet, len(rows) + 6,
        "Every number here comes from the temporal test window only. "
        "assumed_uplift values throughout the workbook are [Guessing] — priors standing in "
        "for evidence that does not exist yet. The experiment sheet is how they get replaced.",
    )
    return sheet


def _sheet_snapshot(workbook, decided):
    sheet = workbook.create_sheet("01 Customer snapshot")
    _title(
        sheet,
        "One row per customer, per point in time",
        "This is what a million transaction rows become before any modelling happens.",
    )

    columns = [
        "customer_id", "snapshot_date", "recency_days", "interpurchase_mean_days",
        "recency_ratio", "frequency", "monetary", "aov", "distinct_products",
        "revenue_trend", "return_rate", "lifecycle", "purchase_probability",
        "purchased_again_30d",
    ]
    sample = (
        decided.sort_values("monetary", ascending=False)
        .drop_duplicates("customer_id")
        .head(20)[columns]
        .copy()
    )

    last = _write_table(
        sheet, sample, start_row=4,
        number_formats={
            "monetary": "#,##0",
            "aov": "#,##0",
            "recency_ratio": "0.00",
            "revenue_trend": "0.00",
            "return_rate": "0.0%",
            "purchase_probability": "0.0%",
            "interpurchase_mean_days": "#,##0",
        },
    )

    _note(
        sheet, last + 2,
        "recency_ratio is the feature worth pointing at: days since last order divided by that "
        "customer's OWN average gap. Above 1.0 they are overdue by their own standard. "
        "Raw recency cannot make that distinction and treats a weekly buyer and a quarterly "
        "buyer identically.",
    )
    return sheet


def _sheet_scorecard(workbook, baseline_scores, model_scores, test_metrics,
                     reliability, stability_frame, importance,
                     observed_rate, mean_balanced, mean_unweighted):
    sheet = workbook.create_sheet("02 Model scorecard")
    _title(sheet, "Does the model work, and can its numbers be trusted?")

    comparison = pd.concat(
        [
            pd.DataFrame(baseline_scores).assign(type="simple rule"),
            pd.DataFrame(model_scores).assign(type="model"),
        ]
    ).sort_values("pr_auc", ascending=False)[["type", "model", "pr_auc"]]

    row = sheet.cell(row=4, column=1, value="BASELINES VS MODELS (validation)")
    row.font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    row.fill = SECTION_FILL
    last = _write_table(sheet, comparison, start_row=5,
                        number_formats={"pr_auc": "0.0000"})

    header = sheet.cell(row=last + 2, column=1, value="TEST PERFORMANCE")
    header.font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    header.fill = SECTION_FILL
    metrics_frame = pd.DataFrame(
        [{"metric": k, "value": v} for k, v in test_metrics.items()]
    )
    last = _write_table(sheet, metrics_frame, start_row=last + 3,
                        number_formats={"value": "0.0000"})

    header = sheet.cell(row=last + 2, column=1, value="CALIBRATION BY DECILE")
    header.font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    header.fill = SECTION_FILL
    last = _write_table(
        sheet, reliability.round(4), start_row=last + 3,
        number_formats={"mean_predicted": "0.0%", "observed_rate": "0.0%", "gap": "0.0%"},
    )

    header = sheet.cell(row=last + 2, column=1, value="STABILITY BY TEST PERIOD")
    header.font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    header.fill = SECTION_FILL
    stability_display = stability_frame.copy()
    for column in stability_display.select_dtypes(include=[np.number]).columns:
        stability_display[column] = stability_display[column].round(4)
    last = _write_table(sheet, stability_display, start_row=last + 3)

    header = sheet.cell(row=last + 2, column=1, value="TOP FEATURES")
    header.font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    header.fill = SECTION_FILL
    last = _write_table(sheet, importance.head(10).round(5), start_row=last + 3)

    last = _note(
        sheet, last + 2,
        f"class_weight='balanced' predicts {mean_balanced:.1%} average probability against an "
        f"observed {observed_rate:.1%}. Unweighted predicts {mean_unweighted:.1%}. Ranking is "
        "barely affected either way — but the decision layer multiplies probability by order "
        "value, so an inflated probability produces an inflated business case.",
    )
    sheet.freeze_panes = None
    return sheet


def _sheet_states(workbook, decided):
    sheet = workbook.create_sheet("03 Customer states")
    _title(
        sheet,
        "The segmentation",
        "State is a description of the situation, not an instruction. Actions come next.",
    )

    profile = (
        decided.groupby("customer_state")
        .agg(
            customers=("customer_id", "size"),
            share=("customer_id", lambda s: len(s) / len(decided)),
            mean_probability=("purchase_probability", "mean"),
            actual_repeat_rate=("purchased_again_30d", "mean"),
            median_recency_days=("recency_days", "median"),
            median_recency_ratio=("recency_ratio", "median"),
            median_monetary=("monetary", "median"),
            median_aov=("aov", "median"),
            median_return_rate=("return_rate", "median"),
        )
        .sort_values("customers", ascending=False)
        .reset_index()
    )

    definitions = {
        "SERVICE_RISK": "Return rate 40%+ with 2+ return events. Outranks every marketing state.",
        "NEW_UNPROVEN": "Exactly one order. Not churned — unmeasured. No rhythm to compare against.",
        "ACTIVE_VALUABLE": "Buying on their own normal rhythm, top 30% by historic spend.",
        "ACTIVE_STANDARD": "Buying on their own normal rhythm, everyone else.",
        "COOLING_VALUABLE": "1-2.5x their own average gap, high value. Slowing, not stopped.",
        "COOLING_STANDARD": "1-2.5x their own average gap, standard value.",
        "LAPSING_VALUABLE": "Beyond 2.5x their gap, high value. The segment worth real money.",
        "LAPSING_STANDARD": "Beyond 2.5x their gap, standard value.",
        "DORMANT": "Over 180 days silent and below the value threshold. Cheap presence only.",
    }
    profile["definition"] = profile["customer_state"].map(definitions)

    last = _write_table(
        sheet, profile, start_row=4,
        number_formats={
            "share": "0.0%",
            "mean_probability": "0.0%",
            "actual_repeat_rate": "0.0%",
            "median_recency_ratio": "0.00",
            "median_monetary": "#,##0",
            "median_aov": "#,##0",
            "median_return_rate": "0.0%",
        },
        width_cap=70,
    )

    last = _note(
        sheet, last + 2,
        "Compare mean_probability against actual_repeat_rate column by column. Close agreement "
        "inside every state is the calibration claim holding up where it matters — at segment "
        "level, which is the level campaigns are planned at.",
    )
    _note(
        sheet, last + 1,
        "LAPSING_VALUABLE is the segment that justifies the whole exercise: a low score and a high "
        "historic value. A pure prediction model ranks them alongside a low-value customer with "
        "the same score. They are not the same customer and should not get the same treatment.",
    )
    return sheet


def _sheet_actions(workbook):
    sheet = workbook.create_sheet("04 Action catalogue")
    _title(
        sheet,
        "Every action, priced by what it costs",
        "Described by cost and mechanism rather than by name — which is what makes them comparable.",
    )

    frame = pd.DataFrame(
        [
            {
                "action": a.key,
                "label": a.label,
                "channel": a.channel,
                "contact_cost": a.contact_cost,
                "incentive_rate": a.incentive_rate,
                "assumed_uplift": a.assumed_uplift,
                "eligible_states": ", ".join(a.eligible_states) or "any",
                "objective": a.objective,
            }
            for a in ACTION_CATALOGUE.values()
        ]
    )

    last = _write_table(
        sheet, frame, start_row=4,
        number_formats={
            "contact_cost": "#,##0.00",
            "incentive_rate": "0.0%",
            "assumed_uplift": "0.0%",
        },
        width_cap=60,
    )

    for i in range(len(frame)):
        cell = sheet.cell(row=5 + i, column=6)
        cell.fill = ASSUMPTION_FILL
        cell.font = Font(name=FONT, size=10, bold=True, color=BLUE)

    _note(
        sheet, last + 2,
        "assumed_uplift is highlighted because it is the only guessed number in the model. "
        "[Guessing] — a prior standing in for evidence that does not exist yet. Every completed "
        "experiment replaces one of these cells with a measurement.",
    )
    return sheet


def _sheet_breakeven(workbook, median_aov):
    """Live formulas. This is the sheet to demo on camera."""
    sheet = workbook.create_sheet("05 Break-even calculator")
    _title(
        sheet,
        "Break-even uplift calculator",
        "Change the blue cells. Everything below recalculates.",
    )

    sheet["A4"] = "INPUTS"
    sheet["A4"].font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    sheet["A4"].fill = SECTION_FILL
    sheet["B4"].fill = SECTION_FILL

    inputs = [
        ("Average order value", round(median_aov, 2), "#,##0.00"),
        ("Gross margin rate", GROSS_MARGIN_RATE, "0.0%"),
        ("Contact cost per customer", 0.15, "#,##0.00"),
        ("Incentive rate (share of AOV)", 0.15, "0.0%"),
        ("Uplift we assume the action creates", 0.045, "0.0%"),
    ]

    for i, (label, value, fmt) in enumerate(inputs, start=5):
        sheet.cell(row=i, column=1, value=label).font = BODY_FONT
        cell = sheet.cell(row=i, column=2, value=value)
        cell.font = INPUT_FONT
        cell.fill = ASSUMPTION_FILL
        cell.number_format = fmt
        cell.border = BOX

    # B5 aov, B6 margin, B7 contact cost, B8 incentive rate, B9 assumed uplift

    sheet["A11"] = "THE EQUATION"
    sheet["A11"].font = Font(name=FONT, bold=True, color=BURGUNDY, size=11)
    sheet["A11"].fill = SECTION_FILL
    sheet["A12"] = "d_breakeven = (contact_cost + P x AOV x incentive_rate) / (AOV x (margin - incentive_rate))"
    sheet["A12"].font = Font(name=FONT, size=10, italic=True, color=INK)

    sheet["A13"] = "Margin minus incentive (must stay positive)"
    sheet["A13"].font = BODY_FONT
    sheet["B13"] = "=B6-B8"
    sheet["B13"].font = FORMULA_FONT
    sheet["B13"].number_format = "0.0%"
    sheet["C13"] = '=IF(B13<=0,"IMPOSSIBLE - the discount exceeds the margin","OK")'
    sheet["C13"].font = FORMULA_FONT

    header_row = 16
    headers = [
        "P(purchase anyway)",
        "Subsidy paid to organic buyers",
        "Required uplift to break even",
        "Assumed uplift",
        "Margin of safety",
        "Verdict",
        "Expected net margin per customer",
    ]
    for j, text in enumerate(headers):
        cell = sheet.cell(row=header_row, column=1 + j, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = BOX

    probabilities = [0.02, 0.05, 0.10, 0.15, 0.20, 0.30, 0.40,
                     0.50, 0.60, 0.70, 0.80, 0.90]

    for i, p in enumerate(probabilities):
        r = header_row + 1 + i

        sheet.cell(row=r, column=1, value=p).number_format = "0%"
        sheet.cell(row=r, column=1).font = FORMULA_FONT

        sheet.cell(row=r, column=2, value=f"=A{r}*$B$5*$B$8").number_format = "#,##0.00"
        sheet.cell(row=r, column=3, value=(
            f"=IF($B$13<=0,\"\",($B$7+A{r}*$B$5*$B$8)/($B$5*$B$13))"
        )).number_format = "0.00%"
        sheet.cell(row=r, column=4, value="=$B$9").number_format = "0.00%"
        sheet.cell(row=r, column=5, value=(
            f"=IFERROR($B$9/C{r},\"\")"
        )).number_format = "0.0x"
        sheet.cell(row=r, column=6, value=(
            f'=IF(C{r}="","impossible",IF($B$9>=C{r},"worth it","wasteful"))'
        ))
        sheet.cell(row=r, column=7, value=(
            f"=$B$9*$B$5*$B$6-$B$7-(A{r}+$B$9)*$B$5*$B$8"
        )).number_format = "#,##0.00;(#,##0.00);-"

        for c in range(1, 8):
            cell = sheet.cell(row=r, column=c)
            cell.border = BOX
            if cell.font is None or cell.font.color is None:
                cell.font = FORMULA_FONT
            else:
                cell.font = FORMULA_FONT
            if i % 2 == 1:
                cell.fill = BAND_FILL

    for j, width in enumerate([26, 26, 24, 16, 16, 14, 24]):
        sheet.column_dimensions[get_column_letter(1 + j)].width = width

    note_row = header_row + len(probabilities) + 2
    note_row = _note(
        sheet, note_row,
        "Read the third column down the page. With any incentive at all, the required uplift RISES "
        "with P — because the discount is paid to everyone who converts, including the customers "
        "who were buying anyway. That is the deadweight cost, and it is why 'P > 0.5 -> send a "
        "discount' targets exactly the wrong people.",
    )
    note_row = _note(
        sheet, note_row,
        "Set the incentive rate to 0% and watch the whole column collapse to near zero: a free "
        "owned-channel message clears its bar almost everywhere. Set the incentive above the "
        "margin and cell C13 tells you no uplift can ever pay for it.",
    )
    _note(
        sheet, note_row,
        "AOV defaults to the median of the scored population. Margin is an assumption supplied by "
        "the user, not derived from the data — Online Retail II carries no cost information.",
    )
    return sheet


def _sheet_state_action(workbook, decided):
    sheet = workbook.create_sheet("06 State x action")
    _title(
        sheet,
        "Which state receives which action",
        "The state decides the action. The score alone never does.",
    )

    matrix = pd.crosstab(
        decided["customer_state"], decided["recommended_action"]
    ).reset_index()

    last = _write_table(sheet, matrix, start_row=4, width_cap=24)

    _note(
        sheet, last + 2,
        "Three cells worth pausing on: ACTIVE_VALUABLE receives loyalty and cross-sell, never a "
        "discount. LAPSING_VALUABLE receives the expensive win-back that a pure propensity ranking "
        "would never justify. SERVICE_RISK receives no campaign at all — it is routed to operations, "
        "because a discount aimed at a product-fit problem makes the returns worse.",
    )
    return sheet


def _sheet_priority(workbook, allocated):
    sheet = workbook.create_sheet("07 Priority customers")
    _title(
        sheet,
        "The list a marketer would actually work",
        "Each row is a complete argument, not a score asking to be trusted.",
    )

    columns = [
        "customer_id", "customer_state", "purchase_probability", "recency_days",
        "recency_ratio", "frequency", "monetary", "aov", "recommended_action",
        "required_uplift", "assumed_uplift", "margin_of_safety",
        "expected_net_margin", "priority_band", "experiment_arm",
    ]
    top = (
        allocated[allocated["contacted"]]
        .sort_values("priority_score", ascending=False)
        .head(30)[columns]
        .copy()
    )

    last = _write_table(
        sheet, top, start_row=4,
        number_formats={
            "purchase_probability": "0.0%",
            "recency_ratio": "0.00",
            "monetary": "#,##0",
            "aov": "#,##0",
            "required_uplift": "0.00%",
            "assumed_uplift": "0.0%",
            "margin_of_safety": "0.0x",
            "expected_net_margin": "#,##0.00",
        },
        width_cap=26,
    )

    _note(
        sheet, last + 2,
        "Read one row aloud and it becomes a sentence: this customer is LAPSING_VALUABLE, scores "
        "31%, has spent heavily, and is 2.8x past their own normal gap. The recommended win-back "
        "needs 0.9pp of incremental conversion to break even; we assume 7pp. That is a decision a "
        "marketer can accept, reject or argue with.",
    )
    return sheet


def _sheet_budget(workbook, budget):
    sheet = workbook.create_sheet("08 Budget")
    _title(
        sheet,
        "What this cycle costs, and what it must deliver",
        "median_required_uplift_pp is the column to argue about in a planning meeting.",
    )

    if budget.empty:
        sheet["A4"] = "No customers selected for contact."
        return sheet

    frame = budget.copy()
    last = _write_table(
        sheet, frame, start_row=4,
        number_formats={
            "fixed_cost": "#,##0.00",
            "expected_incentive_cost": "#,##0.00",
            "total_cost": "#,##0.00",
            "median_required_uplift_pp": "0.00",
            "expected_net_margin_if_assumption_holds": "#,##0.00",
        },
        width_cap=32,
    )

    total_row = last + 1
    sheet.cell(row=total_row, column=1, value="TOTAL").font = Font(
        name=FONT, bold=True, color=BURGUNDY, size=10
    )
    for column_name, letter in [
        ("customers", None), ("fixed_cost", None),
        ("expected_incentive_cost", None), ("total_cost", None),
    ]:
        if column_name not in frame.columns:
            continue
        index = list(frame.columns).index(column_name) + 1
        letter = get_column_letter(index)
        cell = sheet.cell(row=total_row, column=index,
                          value=f"=SUM({letter}5:{letter}{last})")
        cell.font = Font(name=FONT, bold=True, size=10, color=INK)
        cell.number_format = "#,##0" if column_name == "customers" else "#,##0.00"
        cell.border = BOX

    _note(
        sheet, total_row + 2,
        "Notice the split between delivery cost and incentive cost. Owned channels reach thousands "
        "for almost nothing; the money concentrates in the few cells where margin is handed over. "
        "That concentration is exactly what the break-even test is there to police.",
    )
    return sheet


def _sheet_experiment(workbook, design):
    sheet = workbook.create_sheet("09 Experiment design")
    _title(
        sheet,
        "Control cells, and what they can actually detect",
        f"Every treated cell reserves {CONTROL_HOLDOUT_SHARE:.0%} as control, assigned by deterministic hash.",
    )

    if design.empty:
        sheet["A4"] = "No experiment cells."
        return sheet

    frame = design.copy()
    frame["assumed_uplift_pp"] = [
        ACTION_CATALOGUE[a].assumed_uplift * 100 if a in ACTION_CATALOGUE else np.nan
        for a in frame["action"]
    ]
    frame["can_detect_it"] = np.where(
        frame["detectable_uplift_pp"] <= frame["assumed_uplift_pp"], "yes", "no"
    )

    last = _write_table(
        sheet, frame, start_row=4,
        number_formats={
            "baseline_rate": "0.0%",
            "detectable_uplift_pp": "0.00",
            "assumed_uplift_pp": "0.00",
        },
        width_cap=30,
    )

    last = _note(
        sheet, last + 2,
        "Any cell where can_detect_it says 'no' is too small to prove the effect it is assuming. "
        "Running it produces a number nobody can interpret either way. Fix it by pooling several "
        "cycles, enlarging the cell, or saying out loud that the result is directional only.",
    )
    _note(
        sheet, last + 1,
        "Next cycle the measurement is one subtraction: repeat_rate(TREATMENT) minus "
        "repeat_rate(CONTROL), per cell. That number replaces the assumed uplift on sheet 04, and "
        "becomes the training data for an uplift model — where the question changes from 'who will "
        "buy' to 'who will buy BECAUSE we acted'.",
    )
    return sheet



def _sheet_playbook(workbook, decided):
    """The segmentation as an instruction, not a count."""
    sheet = workbook.create_sheet("10 Playbook")
    _title(
        sheet,
        "What each segment actually gets",
        "The matrix says how many. This says what to do, what not to do, and how it is judged.",
    )

    counts = decided["customer_state"].value_counts()
    top_action = (
        decided.groupby("customer_state")["recommended_action"]
        .agg(lambda s: s.value_counts().index[0])
    )
    stats = decided.groupby("customer_state").agg(
        mean_probability=("purchase_probability", "mean"),
        median_monetary=("monetary", "median"),
        median_recency_days=("recency_days", "median"),
    )

    order = [
        "SERVICE_RISK", "NEW_UNPROVEN", "ACTIVE_VALUABLE", "ACTIVE_STANDARD",
        "COOLING_VALUABLE", "COOLING_STANDARD", "LAPSING_VALUABLE",
        "LAPSING_STANDARD", "DORMANT",
    ]
    order = [s for s in order if s in counts.index]

    rows = []
    for state in order:
        entry = STATE_PLAYBOOK[state]
        rows.append(
            {
                "customer_state": state,
                "customers": int(counts[state]),
                "share": counts[state] / len(decided),
                "avg_propensity": float(stats.loc[state, "mean_probability"]),
                "median_spend": float(stats.loc[state, "median_monetary"]),
                "median_days_silent": float(stats.loc[state, "median_recency_days"]),
                "who_they_are": entry["who"],
                "objective": entry["objective"],
                "engine_action": top_action[state],
                "what_to_do": "\n".join(f"{i}. {step}" for i, step in enumerate(entry["do"], 1)),
                "what_to_avoid": entry["avoid"],
                "how_it_is_judged": entry["measure"],
            }
        )

    frame = pd.DataFrame(rows)

    last = _write_table(
        sheet, frame, start_row=4,
        number_formats={
            "share": "0.0%",
            "avg_propensity": "0.0%",
            "median_spend": "#,##0",
            "median_days_silent": "#,##0",
        },
        width_cap=64,
    )

    for i in range(len(frame)):
        row = 5 + i
        sheet.row_dimensions[row].height = 92
        for column in range(1, len(frame.columns) + 1):
            sheet.cell(row=row, column=column).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
        sheet.cell(row=row, column=1).font = Font(name=FONT, size=10, bold=True, color=BURGUNDY)
        avoid_col = list(frame.columns).index("what_to_avoid") + 1
        judged_col = list(frame.columns).index("how_it_is_judged") + 1
        sheet.cell(row=row, column=avoid_col).font = Font(name=FONT, size=10, color=ORANGE)
        sheet.cell(row=row, column=judged_col).font = Font(name=FONT, size=10, bold=True, color=TEAL)

    for name, width in [
        ("who_they_are", 34), ("objective", 34), ("what_to_do", 62),
        ("what_to_avoid", 46), ("how_it_is_judged", 34),
    ]:
        index = list(frame.columns).index(name) + 1
        sheet.column_dimensions[get_column_letter(index)].width = width

    last = _note(
        sheet, last + 2,
        "how_it_is_judged is the column people skip, and it is where most segmentation work quietly "
        "fails. Judge a loyalty programme on conversion rate and it looks useless — those customers "
        "were converting anyway. Judge a service review on orders and you will cancel it. Each state "
        "carries the metric its objective implies.",
    )
    _note(
        sheet, last + 1,
        "engine_action is what the break-even maths selected. what_to_do is the human execution "
        "behind that label. The two are separate on purpose: change the margin assumption and the "
        "action can move while the play stays the same.",
    )
    return sheet


# ==========================================================================
# BUILD
# ==========================================================================

def build_workbook(path, decided, allocated, budget, design, reliability,
                   stability_frame, importance, baseline_scores, model_scores,
                   test_metrics, best_name, median_aov, observed_rate,
                   mean_balanced, mean_unweighted):
    path = Path(path)

    workbook = Workbook()
    workbook.remove(workbook.active)

    _sheet_legend(workbook, best_name, test_metrics, observed_rate,
                  mean_balanced, mean_unweighted, median_aov)
    _sheet_snapshot(workbook, decided)
    _sheet_scorecard(workbook, baseline_scores, model_scores, test_metrics,
                     reliability, stability_frame, importance,
                     observed_rate, mean_balanced, mean_unweighted)
    _sheet_states(workbook, decided)
    _sheet_actions(workbook)
    _sheet_breakeven(workbook, median_aov)
    _sheet_state_action(workbook, decided)
    _sheet_priority(workbook, allocated)
    _sheet_budget(workbook, budget)
    _sheet_experiment(workbook, design)
    _sheet_playbook(workbook, decided)

    workbook.save(path)
    return path